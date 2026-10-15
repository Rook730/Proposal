"""Import one proposal sheet (e.g. "26005 Proposal") from the Excel workbook.

Built for the KS Proposal Template layout: each section starts with a row
whose column A holds the table name and column B "=Items!C..", each line is a
row with a cost formula in K and its item code two rows below in B.

    python -m proposal_builder.import_proposal Proposal_Template4.xlsx "26005 Proposal" [--data DIR]
"""

from __future__ import annotations

import argparse
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from .importer import _dec, _text
from .models import Catalog, Performer, PriceSource, Settings, Variable
from .proposal import (
    Alternate,
    Line,
    LineTier,
    Project,
    ProjectStatus,
    Revision,
    default_blocks,
    line_from_item,
)
from .store import DataStore, ProjectStore, resolve_data_dir

CW_PREFIX = "Performed by Convenient Woodworking."
ROW_RE = re.compile(r"K(\d+)(?::K(\d+))?")


def _rows_in(formula: str) -> set[int]:
    """'=K268+SUM(K274:K315)' -> {268, 274, ..., 315}"""
    rows: set[int] = set()
    for a, b in ROW_RE.findall(formula or ""):
        rows |= set(range(int(a), int(b or a) + 1))
    return rows


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    for fmt in ("%Y.%m.%d", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return date.today()


def _rev_str(v) -> str:
    d = _dec(v)
    if d is None:
        return "1.0"
    return f"{d:.1f}" if d == d.to_integral_value() else f"{d.normalize():f}"


def _extract_variables(text: str) -> list[Variable]:
    found = []
    pats = [
        ("meeting_dates", r"at their home on (.+?)\.", "Site visit dates quoted in the cover letter"),
        ("working_time", r"Total working time is approximately (.+?),", "Duration quoted in letter and schedule"),
        ("hood_cfm", r"(\d+)\s*CFM", "Range hood airflow cap"),
    ]
    for name, pat, desc in pats:
        m = re.search(pat, text)
        if m:
            found.append(Variable(name=name, value=m.group(1).strip(), description=desc))
    return found


def import_proposal_sheet(xlsx: Path, sheet: str, catalog: Catalog, settings: Settings,
                          boilerplate_blocks) -> tuple[Project, Revision]:
    wf = openpyxl.load_workbook(xlsx)[sheet]
    wv = openpyxl.load_workbook(xlsx, data_only=True)[sheet]

    def f(cell: str):
        return wf[cell].value

    def v(cell: str):
        return wv[cell].value

    # header (labels in H, values in I; client block in B/C)
    hdr = {}
    for r in range(1, 20):
        label = _text(wf.cell(r, 8).value).rstrip(":")
        if label:
            hdr[label] = wv.cell(r, 9).value
        label_b = _text(wf.cell(r, 2).value).rstrip(":")
        if label_b in ("Client", "Address", "Project"):
            hdr[label_b] = _text(wv.cell(r, 3).value)
    number = _text(hdr.get("Proposal #")) or sheet.split()[0]
    project = Project(number=number, client_name=hdr.get("Client", ""), site_address=hdr.get("Address", ""),
                      project_name=hdr.get("Project", ""), status=ProjectStatus.PRESENTED)
    rev = Revision(rev=_rev_str(hdr.get("Proposal Revision")), proposal_date=_parse_date(hdr.get("Proposal Date")),
                   blocks=default_blocks(boilerplate_blocks))

    by_key = {s.key: s for s in catalog.sections}
    by_name = {s.name.lower(): s for s in catalog.sections}

    current = None
    row_to_line: dict[int, Line] = {}
    alt_with: set[int] = set()
    alt_excl: set[int] = set()
    for r in range(1, wf.max_row + 1):
        a = _text(wf.cell(r, 1).value)
        b_formula = _text(wf.cell(r, 2).value)
        if a and b_formula.startswith("=Items!"):
            sec = by_key.get(a) or by_name.get(_text(wv.cell(r, 2).value).lower())
            current = sec
            if sec and sec.number not in rev.sections:
                rev.sections.append(sec.number)
            continue
        label = _text(wv.cell(r, 2).value)
        if "(With alternate items)" in label:
            alt_with |= _rows_in(_text(wf.cell(r, 10).value))
        elif "(Excluding alternate items)" in label:
            alt_excl |= _rows_in(_text(wf.cell(r, 10).value))

        k = _text(wf.cell(r, 11).value)
        code_v = wf.cell(r + 2, 2).value
        if current is None or not (k.startswith("=") and "*" in k):
            continue
        try:
            code = int(code_v)
        except (TypeError, ValueError):
            continue
        row_to_line[r] = _build_line(r, code, current.number, wf, wv, catalog, rev, k)

    # alternates: rows only in the "with" sum are added, rows only in the "excluding" sum are removed
    added = sorted(r for r in alt_with - alt_excl if r in row_to_line)
    removed = sorted(r for r in alt_excl - alt_with if r in row_to_line)
    if added or removed:
        alt = Alternate(number=1, name="Alternate 1",
                        removes=[row_to_line[r].id for r in removed])
        for r in added:
            row_to_line[r].alternate = 1
        first = row_to_line[added[0]] if added else None
        if first is not None:
            alt.description = first.notes or first.title
        rev.alternates.append(alt)

    all_text = "\n".join(_text(wv.cell(r, 2).value) for r in range(1, wv.max_row + 1))
    rev.variables = _extract_variables(all_text)

    # boilerplate: match the sheet. A library block is on when its sub-heading appears in the sheet;
    # the generic "Room Availability" block takes the sheet's "<Room> Availability" title.
    titles = {_text(wv.cell(r, 2).value) for r in range(1, wv.max_row + 1)}
    for b in rev.blocks:
        if not b.title:
            continue
        if b.title == "Room Availability":
            found = next((t for t in titles if t.endswith(" Availability") and len(t) < 40), None)
            if found:
                b.title, b.on = found, True
            continue
        if b.heading in ("Payment Schedule", "Payment Terms & Contract Cancellation",
                         "Scope Clarifications & Exclusions"):
            b.on = b.title in titles
    return project, rev


def _build_line(r: int, code: int, section: int, wf, wv, catalog: Catalog, rev: Revision, k_formula: str) -> Line:
    item = catalog.item(code)
    title = _text(wv.cell(r, 2).value)
    desc_f = wf.cell(r, 4).value
    desc_v = _text(wv.cell(r, 4).value)
    unit = _text(wv.cell(r, 8).value) or "LS"
    qty = _dec(wv.cell(r, 9).value) or Decimal(0)
    price_f = wf.cell(r, 10).value
    price_v = _dec(wv.cell(r, 10).value)
    price_typed = not (isinstance(price_f, str) and price_f.startswith("="))
    desc_typed = not (isinstance(desc_f, str) and str(desc_f).startswith("="))
    sec = item.section if item else section

    if item:
        ln = line_from_item(item, catalog, rev.new_line_id(sec))
    else:
        ln = Line(id=rev.new_line_id(sec), section=sec, code=code, title=title or f"Item {code}")
    ln.title = title or ln.title
    ln.unit = unit
    ln.qty = qty

    if desc_typed and desc_v:
        notes = desc_v
        if notes.startswith(CW_PREFIX):
            ln.performer = Performer.CW
            notes = notes[len(CW_PREFIX):].strip()
        ln.notes = notes
        ln.print_general = False
        if (price_v or 0) == 0 and "purchased directly by client" in desc_v.lower():
            ln.owner_furnished = True
            ln.performer = Performer.CLIENT

    if "Subtotal" in k_formula:
        ln.price_basis = PriceSource.PERCENT
        ln.unit_price = price_v
        ln.qty = Decimal(1)
    elif price_typed:
        _apply_typed_price(ln, price_v, catalog)
    elif price_v is not None and ln.unit_price != price_v:
        ln.unit_price = price_v

    rev.add_line(ln)
    return ln


def _apply_typed_price(ln: Line, price: Decimal | None, catalog: Catalog) -> None:
    if price is None:
        return
    if ln.price_basis is PriceSource.RATE:
        if ln.unit_price == price:
            return
        same_rate = [r for r in catalog.rates if r.rate == price and r.unit == ln.unit]
        preferred = "cw" if ln.performer is Performer.CW else None
        same = next((r for r in same_rate if r.key == preferred), same_rate[0] if same_rate else None)
        if same:
            ln.rate_key, ln.unit_price = same.key, same.rate
            return
    if ln.price_basis is PriceSource.ALLOWANCE:
        ln.tier = LineTier.CUSTOM
        ln.unit_price = price
        return
    ln.price_basis = PriceSource.FIXED
    ln.unit_price = price


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("sheet")
    ap.add_argument("--data", type=Path, help="data folder (default: saved setting)")
    args = ap.parse_args(argv)
    ds = DataStore(args.data or resolve_data_dir())
    if not ds.is_initialized():
        ds.initialize_from_seed()
    project, rev = import_proposal_sheet(args.xlsx, args.sheet, ds.load_catalog(), ds.load_settings(),
                                         ds.load_boilerplate().blocks)
    folder = ProjectStore(ds).create(project, rev)
    print(f"Imported {project.number} rev {rev.rev}: {len(rev.lines)} lines, "
          f"{len(rev.alternates)} alternate(s) -> {folder}")


if __name__ == "__main__":
    main()
