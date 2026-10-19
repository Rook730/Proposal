"""One-time import from the Excel workbook (Proposal_Template4.xlsx).

Reads the Items, Allowances and Definitions tabs plus the boilerplate text of
the "26005 Proposal" tab, applies the clean-up listed in docs/requirements.md
section 13, and writes settings.json, catalog.json and boilerplate.json.

    python -m proposal_builder.importer Proposal_Template4.xlsx --out <folder>
"""

from __future__ import annotations

import argparse
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from .models import (
    AllowanceRange,
    Boilerplate,
    BoilerplateBlock,
    Catalog,
    CatalogItem,
    Company,
    Performer,
    PriceSource,
    Rate,
    Section,
    Settings,
    StockClass,
    Tier,
    Unit,
)
from .store import _atomic_write, dump_model

LAST_RENOVATION_CODE = 1999          # 2000+ is lamp restoration: dropped
PROPOSAL_SHEET = "26005 Proposal"

# Section keys, taken from the table names used in the workbook.
SECTION_KEYS = {
    100: "GeneralConditions", 200: "ManagementSupervision", 300: "SiteProtection",
    400: "TempUtilities", 500: "Demo", 600: "MEP", 700: "Insulation",
    800: "IntConst", 900: "Millwork", 1000: "Floor", 1100: "InteriorFinish",
    1200: "Closeout", 1300: "Misc", 1400: "Overhead",
}

RATES = [
    Rate(key="pm", name="Project management / client meetings", rate=Decimal("75")),
    Rate(key="admin", name="Pre-construction / administration / inspections", rate=Decimal("60")),
    Rate(key="protection", name="Protective services labor", rate=Decimal("45")),
    Rate(key="cleanup", name="Cleanup labor", rate=Decimal("35")),
    Rate(key="general", name="General labor (demo, doors)", rate=Decimal("50")),
    Rate(key="cw", name="Convenient Woodworking trade labor", rate=Decimal("75")),
    Rate(key="licensed", name="Licensed trade labor (electrical, plumbing)", rate=Decimal("95")),
    Rate(key="fixtures", name="Fixtures installation labor", rate=Decimal("45")),
]

ITEM_RATE = {
    210: "pm", 250: "pm", 1210: "pm", 1230: "pm",
    220: "admin", 230: "admin", 240: "admin",
    310: "protection", 330: "cleanup", 340: "cleanup",
    510: "general", 820: "general",
    616: "cw", 860: "cw", 920: "cw", 931: "cw", 941: "cw", 1040: "cw",
    626: "licensed", 636: "licensed",
    1150: "fixtures",
}

PERCENT_ITEMS = {130, 1310, 1320, 1410, 1420}
CREDIT_IF_UNUSED = {1310, 1320, 1110}

PERFORMER = {
    **{c: Performer.KEPT_SPEC for c in (210, 220, 230, 240, 250, 1210, 1220, 1230, 1240)},
    **{c: Performer.CW for c in (310, 330, 340, 510, 616, 820, 860, 920, 931, 941, 1040, 1150)},
    **{c: Performer.SUB for c in (615, 625, 635, 645, 710, 810, 840, 950, 1020, 1110, 1160)},
}

DESCRIPTION_FIXES = {
    450: "Allowance for utility connection fees associated with the project.",
    1320: ("5% allowance for project management of change orders that may arise during the project: scoping, "
           "pricing, documentation and rescheduling. Time spent is recorded on each change order and drawn from "
           "this allowance at the project management rate.\n*Unused allowance is credited on the final invoice."),
}

NEW_ITEMS = [
    CatalogItem(code=1121, section=1100, title="Tile material", unit="SF",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of tile material. Price subject to change with final material selection."),
    CatalogItem(code=1141, section=1100, title="Cabinet hardware", unit="EA",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of cabinet hardware. Price subject to change with final selection."),
    CatalogItem(code=1142, section=1100, title="Lighting fixtures", unit="EA",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of lighting fixtures. Price subject to change with final selection."),
    CatalogItem(code=1143, section=1100, title="Plumbing fixtures", unit="EA",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of plumbing fixtures. Price subject to change with final selection."),
    CatalogItem(code=1144, section=1100, title="Glass and mirrors", unit="EA",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of glass and mirrors. Price subject to change with final selection."),
    CatalogItem(code=1171, section=1100, title="Ventilation equipment", unit="EA",
                price_source=PriceSource.ALLOWANCE,
                description="Allowance for the purchase of ventilation equipment. Price subject to change with final selection."),
]

# allowance category -> catalog item code; (category, description-prefix) overrides
ALLOWANCE_LINK = {
    "Electrical": 627, "Flooring": 1010, "Glass": 1144, "Hardware": 1141,
    "Lighting": 1142, "Millwork": 940, "Paint": 1110, "Plumbing Fixtures": 1143,
    "Tile": 1121, "Ventilation": 1171,
}
ALLOWANCE_LINK_OVERRIDE = {
    ("Tile", "Backsplash install"): 1160,
    ("Ventilation", "Duct kit"): 617,
}

STOCK_CLASSES = [
    StockClass(key="trim", name="Trim / millwork", lengths_ft=[Decimal(x) for x in (8, 10, 12, 16, 20)]),
]
STOCK_ITEMS = {940: "trim"}


# ---------------------------------------------------------------- helpers


def _dec(v) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = re.sub(r"[$,\s]", "", str(v))
    if s in ("", "-"):
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _slug(*parts: str) -> str:
    s = "-".join(parts).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


# ---------------------------------------------------------------- catalog


def import_units(wb) -> list[Unit]:
    ws = wb["Definitions"]
    units = []
    for row in ws.iter_rows(min_row=3, min_col=2, max_col=3, values_only=True):
        code, desc = row
        if code:
            units.append(Unit(code=_text(code), description=_text(desc)))
    if not any(u.code == "PR" for u in units):
        units.append(Unit(code="PR", description="Pair"))
    return sorted(units, key=lambda u: u.code)


def import_items(wb) -> tuple[list[Section], list[CatalogItem]]:
    ws = wb["Items"]
    sections: list[Section] = []
    items: list[CatalogItem] = []
    group = ""
    for row in ws.iter_rows(min_col=2, max_col=6, values_only=True):
        code_raw, title, desc, unit, price = row
        try:
            code = int(code_raw)
        except (TypeError, ValueError):
            continue
        if code > LAST_RENOVATION_CODE:
            continue
        title, desc, unit = _text(title), _text(desc), _text(unit)
        if desc == "-":
            desc = ""
        if code % 100 == 0:
            if title:
                sections.append(Section(number=code, name=title, key=SECTION_KEYS.get(code, _slug(title))))
            group = ""
            continue
        if not title:
            continue                                     # empty placeholder code
        if not unit and not desc and price is None:
            group = title                                # sub-heading such as "Electrical"
            continue
        section = code // 100 * 100
        desc = DESCRIPTION_FIXES.get(code, desc)
        price_d = _dec(price)
        item = CatalogItem(
            code=code,
            section=section,
            group=group if section == 600 else "",
            title=title,
            description=desc,
            unit=unit or "LS",
            performer=PERFORMER.get(code, Performer.NONE),
            not_included="not included" in desc.lower() and (price_d or 0) == 0,
            credit_if_unused=code in CREDIT_IF_UNUSED,
            stock_class=STOCK_ITEMS.get(code),
        )
        if code in PERCENT_ITEMS:
            item.price = price_d
            item.price_source = PriceSource.PERCENT
        elif code in ITEM_RATE:
            item.rate_key = ITEM_RATE[code]
            item.price_source = PriceSource.RATE
        else:
            item.price = price_d
        items.append(item)
    return sections, items


def import_allowances(wb) -> list[AllowanceRange]:
    ws = wb["Allowances"]
    out = []
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=7, values_only=True):
        cat, desc, unit, low, high, basis = row
        if not cat or not desc:
            continue
        cat, desc = _text(cat), _text(desc)
        out.append(AllowanceRange(
            id=_slug(cat, desc), category=cat, description=desc, unit=_text(unit),
            low=_dec(low) or Decimal(0), high=_dec(high) or Decimal(0), basis=_text(basis),
        ))
    return out


def link_allowances(items: list[CatalogItem], allowances: list[AllowanceRange]) -> None:
    by_code = {i.code: i for i in items}
    for a in allowances:
        code = ALLOWANCE_LINK.get(a.category)
        for (cat, prefix), c in ALLOWANCE_LINK_OVERRIDE.items():
            if a.category == cat and a.description.startswith(prefix):
                code = c
        if code in by_code:
            by_code[code].allowance_ids.append(a.id)


def build_catalog(wb) -> Catalog:
    units = import_units(wb)
    sections, items = import_items(wb)
    have = {i.code for i in items}
    items += [i.model_copy(deep=True) for i in NEW_ITEMS if i.code not in have]
    items.sort(key=lambda i: i.code)
    allowances = import_allowances(wb)
    link_allowances(items, allowances)
    for i in items:
        if i.price_source is PriceSource.ALLOWANCE:
            i.default_tier = Tier.HIGH
    cat = Catalog(units=units, rates=[r.model_copy() for r in RATES], sections=sections,
                  items=items, allowances=allowances,
                  stock_classes=[s.model_copy(deep=True) for s in STOCK_CLASSES])
    problems = cat.validate_refs()
    if problems:
        raise ValueError("catalog import problems:\n" + "\n".join(problems))
    return cat


# ---------------------------------------------------------------- settings


def build_settings(wb) -> Settings:
    ws = wb[PROPOSAL_SHEET]
    header = _text(ws["B4"].value)
    parts = [p.strip() for p in header.split("|")]
    address = parts[0] if parts else ""
    phone = parts[1].replace("Phone:", "").strip() if len(parts) > 1 else ""
    email = parts[2].replace("Email:", "").strip() if len(parts) > 2 else ""
    return Settings(
        company=Company(legal_name="Convenient Enterprises LLC", dba="Kept Spec Advisors",
                        address=address, phone=phone, email=email, logo_path="logo.png"),
        proposal_year=26, next_proposal_seq=6,
    )


# ---------------------------------------------------------------- boilerplate

# (paragraph prefix, block id, printed title, strip "Title." lead-in from body)
HOW_TO_READ_TITLES = [
    ("The work is organized", "sections", "Sections and line items", False),
    ("Each line item has", "line-items", "Line items", False),
    ("Hourly (HR)", "hourly", "Hourly quantities", False),
    ("Lines described as", "allowances", "Allowances", False),
    ("Who performs the work.", "performer", "Who performs the work", True),
    ("$0 lines.", "zero-lines", "$0 lines", True),
    ("Gray lines.", "gray-lines", "Gray lines", True),
    ("Blue lines.", "blue-lines", "Blue lines", True),
    ("Exclusions.", "exclusions-note", "Exclusions", True),
]

# Generic placeholders replacing values that belong to one proposal.
TEMPLATE_SUBS = [

    ("Total working time is approximately three weeks", "Total working time is approximately {working_time}"),
    ("25% Due Upon Signing", "{deposit_pct} Due Upon Signing"),
    ("Remaining balance due less 5%", "Remaining balance due less {retainage_pct}"),
    ("valid for 30 days", "valid for {validity_days} days"),
    ("Unused contingency is credited on the final invoice.",
     "Unused contingency and unused change order management are credited on the final invoice."),
    ("hood limited to 400 CFM", "hood limited to {hood_cfm} CFM"),
]


# Regex substitutions for the parts of the cover letter that name one client's project.
TEMPLATE_REGEX = [
    (r"itemized breakdown for the .+? at .+?\. This proposal",
     "itemized breakdown for the {project_name} at {site_address}. This proposal"),
    (r"at their home on [^.]+?\.", "at their home on {meeting_dates}."),
]


def _templatize(text: str) -> str:
    for old, new in TEMPLATE_SUBS:
        text = text.replace(old, new)
    for pat, new in TEMPLATE_REGEX:
        text = re.sub(pat, new, text, count=1)
    return text


def _md_lists(text: str) -> str:
    """Excel used "- " bullets already; normalize CRLF and trim."""
    return text.replace("\r\n", "\n").strip()


def _split_how_to_read(text: str) -> list[tuple[str, str, str]]:
    paras: list[str] = []
    for chunk in text.split("\n\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        if paras and paras[-1].endswith(":"):
            paras[-1] += "\n\n" + chunk
            continue
        # "Hourly" and "Lines described" share one \n-separated chunk
        for piece in chunk.split("\n") if chunk.startswith("Hourly") else [chunk]:
            paras.append(piece.strip())
    out = []
    for p in paras:
        for prefix, key, title, strip in HOW_TO_READ_TITLES:
            if p.startswith(prefix):
                body = p[len(prefix):].strip() if strip else p
                out.append((key, title, body))
                break
        else:
            out.append((_slug(p[:30]), "", p))
    return out


def _pairs(ws, first: int, last: int) -> list[tuple[str, str]]:
    """Heading/body pairs from column B between two rows (inclusive)."""
    cells = [(r, _text(ws.cell(r, 2).value)) for r in range(first, last + 1)]
    cells = [(r, v) for r, v in cells if v]
    out, i = [], 0
    while i < len(cells) - 1:
        out.append((cells[i][1], cells[i + 1][1]))
        i += 2
    return out


def _find_row(ws, text: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if _text(ws.cell(r, 2).value) == text:
            return r
    raise ValueError(f"'{text}' not found in {ws.title}")


def build_boilerplate(wb) -> Boilerplate:
    ws = wb[PROPOSAL_SHEET]
    blocks: list[BoilerplateBlock] = []

    def add(heading, id_, title, body, order, default_on=True):
        blocks.append(BoilerplateBlock(id=id_, heading=heading, title=title,
                                       body=_templatize(_md_lists(body)), order=order,
                                       default_on=default_on))

    add("Intro", "cover-letter", "", _text(ws["B16"].value), 10)
    add("Intro", "cost-plus", "", _text(ws["B37"].value), 20)

    for n, (key, title, body) in enumerate(_split_how_to_read(_text(ws["B43"].value))):
        add("How to read this proposal", key, title, body, 10 * (n + 1))
    add("How to read this proposal", "units-table", "Unit definitions", "{units_table}", 25)  # after line-items

    add("Totals", "alternate-note", "", _text(ws["B388"].value), 10)
    add("Signature", "thanks", "", _text(ws["B392"].value), 10)
    add("Signature", "acknowledgement", "", _text(ws["B398"].value), 20)

    ps = _find_row(ws, "Payment Schedule")
    pt = _find_row(ws, "Payment Terms & Contract Cancellation", ps)
    sc = _find_row(ws, "Scope Clarifications & Exclusions", pt)
    groups = [
        ("Payment Schedule", ps + 1, pt - 1),
        ("Payment Terms & Contract Cancellation", pt + 1, sc - 1),
        ("Scope Clarifications & Exclusions", sc + 1, ws.max_row),
    ]
    for heading, a, b in groups:
        for n, (title, body) in enumerate(_pairs(ws, a, b)):
            add(heading, _slug(heading[:12], title), title, body, 10 * (n + 1))

    # Payment schedule blocks open with a short terms line (e.g. "25% Due Upon Signing"): bold it.
    for blk in blocks:
        if blk.heading == "Payment Schedule":
            first, sep, rest = blk.body.partition("\n")
            if sep and len(first) < 60 and not first.startswith(("-", "*")) and not first.endswith("."):
                blk.body = f"**{first}**\n{rest}"

    # Project-specific wording becomes a generic room-availability block, off by default.
    for blk in blocks:
        if blk.title == "Kitchen Availability":
            blk.title = "Room Availability"
            blk.default_on = False

    from .added_terms import apply as add_terms
    add_terms(blocks)

    headings = ["Intro", "How to read this proposal", "Totals", "Signature",
                "Payment Schedule", "Payment Terms & Contract Cancellation",
                "Scope Clarifications & Exclusions"]
    return Boilerplate(headings=headings, blocks=blocks)


# ---------------------------------------------------------------- entry point


def import_workbook(xlsx: Path) -> tuple[Settings, Catalog, Boilerplate]:
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    return build_settings(wb), build_catalog(wb), build_boilerplate(wb)


def write_all(out: Path, settings: Settings, catalog: Catalog, boilerplate: Boilerplate) -> None:
    out.mkdir(parents=True, exist_ok=True)
    _atomic_write(out / "settings.json", dump_model(settings))
    _atomic_write(out / "catalog.json", dump_model(catalog))
    _atomic_write(out / "boilerplate.json", dump_model(boilerplate))


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--out", type=Path, required=True, help="folder to write the JSON files into")
    args = ap.parse_args(argv)
    s, c, b = import_workbook(args.xlsx)
    write_all(args.out, s, c, b)
    print(f"{len(c.sections)} sections, {len(c.items)} items, {len(c.allowances)} allowance ranges, "
          f"{len(c.rates)} rates, {len(b.blocks)} boilerplate blocks -> {args.out}")


if __name__ == "__main__":
    main()
